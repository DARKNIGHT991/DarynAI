import base64
import csv
import io
import zipfile
from pathlib import PurePosixPath


TEXT_EXTENSIONS = {
    ".txt", ".js", ".py", ".html", ".css", ".json", ".md", ".csv", ".log",
    ".ts", ".tsx", ".jsx", ".java", ".kt", ".go", ".rs", ".php", ".rb",
    ".cs", ".cpp", ".c", ".h", ".hpp", ".sql", ".xml", ".yaml", ".yml",
    ".toml", ".ini", ".env", ".gitignore", ".dockerignore",
}
PROJECT_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx", ".xlsx"}
SKIP_DIRS = {
    ".git", "node_modules", "dist", "build", ".next", ".nuxt", ".venv",
    "venv", "__pycache__", ".pytest_cache", ".mypy_cache", "coverage",
}
MAX_ZIP_FILES = 120
MAX_ZIP_FILE_BYTES = 1_000_000
MAX_ZIP_TOTAL_BYTES = 5_000_000


def extract_uploaded_file(file_name: str | None, file_data: str) -> str:
    name = file_name or "uploaded_file"
    lower_name = name.lower()
    raw = base64.b64decode(file_data)

    if lower_name.endswith(".pdf"):
        return extract_pdf(raw)
    if lower_name.endswith(".docx"):
        return extract_docx(raw)
    if lower_name.endswith(".xlsx"):
        return extract_xlsx(raw)
    if lower_name.endswith(".csv"):
        return extract_csv(raw)
    if lower_name.endswith(".zip"):
        return extract_project_zip(raw)

    return decode_text(raw)


def extract_pdf(raw: bytes) -> str:
    try:
        import PyPDF2
    except ImportError as exc:
        raise RuntimeError("PDF support requires PyPDF2 in requirements.txt") from exc

    reader = PyPDF2.PdfReader(io.BytesIO(raw))
    pages = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"--- PDF page {page_number} ---\n{text}")
    return "\n\n".join(pages).strip()


def extract_docx(raw: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX support requires python-docx in requirements.txt") from exc

    document = Document(io.BytesIO(raw))
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"--- DOCX table {table_index} ---\n" + "\n".join(rows))

    return "\n\n".join(parts).strip()


def extract_xlsx(raw: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX support requires openpyxl in requirements.txt") from exc

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        rows = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_index > 200:
                rows.append("... sheet truncated after 200 rows ...")
                break
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                rows.append(" | ".join(values))
        if rows:
            sheets.append(f"--- XLSX sheet: {sheet.title} ---\n" + "\n".join(rows))
    return "\n\n".join(sheets).strip()


def extract_csv(raw: bytes) -> str:
    text = decode_text(raw)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample)
        rows = csv.reader(io.StringIO(text), dialect)
        rendered = []
        for index, row in enumerate(rows, start=1):
            if index > 300:
                rendered.append("... csv truncated after 300 rows ...")
                break
            rendered.append(" | ".join(row))
        return "\n".join(rendered)
    except csv.Error:
        return text


def extract_project_zip(raw: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        infos = [info for info in archive.infolist() if not info.is_dir()]
        safe_infos = [info for info in infos if is_safe_zip_member(info.filename)]
        safe_infos = [info for info in safe_infos if not should_skip_path(info.filename)]

        tree_lines = build_zip_tree(safe_infos)
        content_sections = []
        total_bytes = 0
        included_files = 0

        for info in safe_infos:
            if included_files >= MAX_ZIP_FILES:
                content_sections.append("... project truncated: too many files ...")
                break
            if info.file_size > MAX_ZIP_FILE_BYTES:
                continue
            if total_bytes + info.file_size > MAX_ZIP_TOTAL_BYTES:
                content_sections.append("... project truncated: total readable content limit reached ...")
                break

            ext = file_extension(info.filename)
            if ext not in PROJECT_EXTENSIONS:
                continue

            data = archive.read(info)
            try:
                extracted = extract_member_content(info.filename, data)
            except Exception as exc:
                extracted = f"[Could not read this file: {exc}]"

            if extracted.strip():
                included_files += 1
                total_bytes += info.file_size
                content_sections.append(
                    f"--- FILE: {info.filename} ---\n{extracted[:12000]}"
                )

        return (
            "ZIP PROJECT ANALYSIS CONTEXT\n\n"
            "Project structure:\n"
            f"{chr(10).join(tree_lines)}\n\n"
            "Readable important files:\n\n"
            + "\n\n".join(content_sections)
        ).strip()


def extract_member_content(name: str, raw: bytes) -> str:
    lower_name = name.lower()
    if lower_name.endswith(".pdf"):
        return extract_pdf(raw)
    if lower_name.endswith(".docx"):
        return extract_docx(raw)
    if lower_name.endswith(".xlsx"):
        return extract_xlsx(raw)
    if lower_name.endswith(".csv"):
        return extract_csv(raw)
    return decode_text(raw)


def decode_text(raw: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def is_safe_zip_member(name: str) -> bool:
    normalized = name.replace("\\", "/")
    path = PurePosixPath(normalized)
    return not path.is_absolute() and ".." not in path.parts


def should_skip_path(name: str) -> bool:
    parts = PurePosixPath(name.replace("\\", "/")).parts
    return any(part in SKIP_DIRS for part in parts)


def file_extension(name: str) -> str:
    path = PurePosixPath(name.replace("\\", "/"))
    if path.name in {".env", ".gitignore", ".dockerignore"}:
        return path.name
    return path.suffix.lower()


def build_zip_tree(infos: list[zipfile.ZipInfo]) -> list[str]:
    paths = sorted(info.filename.replace("\\", "/") for info in infos)
    lines = []
    for path in paths[:300]:
        depth = max(0, len(PurePosixPath(path).parts) - 1)
        lines.append(f"{'  ' * depth}- {PurePosixPath(path).name}")
    if len(paths) > 300:
        lines.append("... structure truncated after 300 files ...")
    return lines or ["(empty project archive)"]
